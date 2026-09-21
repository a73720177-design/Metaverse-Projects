from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _shape_text(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_text(shape.shapes)
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            yield shape.text.strip()
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                yield " | ".join(cell.text.strip() for cell in row.cells)


def parse_ppt(path: Path) -> list[tuple[int, str]]:
    presentation = Presentation(path)
    slides: list[tuple[int, str]] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts = list(_shape_text(slide.shapes))
        slides.append((slide_number, "\n".join(parts)))
    return slides
