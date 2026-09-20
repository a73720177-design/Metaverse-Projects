import asyncio
import base64
import json
from pathlib import Path

import fitz
import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from app.integrations import vision
from app.parsers import visual_parser
from app.parsers.ppt_parser import parse_ppt
from app.services.document_service import parse_document
from app.services.rag_service import DocumentContextSelector


@pytest.fixture(autouse=True)
def enable_vision(monkeypatch):
    monkeypatch.setenv("DOCUMENT_VISION_MODE", "ollama")
    monkeypatch.setenv("VLM_BASE_URL", "http://127.0.0.1:11434")
    monkeypatch.setenv("VLM_MODEL", "qwen3-vl:4b-instruct")
    monkeypatch.setenv("VLM_MAX_PAGES", "40")


def make_pdf(path, pages=2):
    with fitz.open() as document:
        for index in range(pages):
            page = document.new_page(width=400, height=300)
            if index == 0:
                page.insert_text((20, 20), "Original text")
        document.save(path)


def test_image_only_pages_are_enriched_and_retrievable(tmp_path, monkeypatch):
    path = tmp_path / "slides.pdf"
    make_pdf(path)
    calls = []

    def describe(self, png, text, page, timeout):
        image = fitz.Pixmap(png)
        assert image.width == 1600 and image.height == 1200
        assert 0 < timeout <= 120
        calls.append((page, text))
        return "태양광 설비의 발전량 차트. 단위 kWh."

    monkeypatch.setattr(vision.OllamaVisionClient, "describe", describe)
    result = parse_document(path, "slides.pdf")
    assert calls == [(1, "Original text"), (2, "")]
    assert [section.index for section in result.sections] == [1, 2]
    assert "Original text" in result.full_text
    assert "AI 해석, 원본 확인 필요" in result.full_text
    selected = DocumentContextSelector().select(result, "태양광 발전량")
    assert selected is not None
    assert "태양광" in selected.full_text


def test_pptx_conversion_preserves_slide_mapping_and_cleans_temp(tmp_path, monkeypatch):
    path = tmp_path / "slides.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    directories = []
    monkeypatch.setattr(visual_parser.shutil, "which", lambda _: "/bin/soffice")

    def convert(command, **kwargs):
        assert "--headless" in command
        assert "pdf:impress_pdf_Export" in command
        assert any(arg.startswith("-env:UserInstallation=file:") for arg in command)
        assert kwargs["timeout"] <= 60
        source = Path(command[-1])
        assert source.read_bytes() == path.read_bytes()
        directories.append(source.parent)
        make_pdf(source.with_suffix(".pdf"))

    monkeypatch.setattr(visual_parser.subprocess, "run", convert)
    monkeypatch.setattr(vision.OllamaVisionClient, "describe", lambda *args: "도식 연결 관계")
    document = parse_document(path, path.name)
    assert [section.index for section in document.sections] == [1, 2]
    assert all(not directory.exists() for directory in directories)


def test_tables_and_grouped_text_are_extracted_without_vision(tmp_path):
    path = tmp_path / "table.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Grouped heading"
    table = slide.shapes.add_table(1, 2, 0, 0, Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Revenue"
    table.cell(0, 1).text = "42"
    deck.save(path)
    assert parse_ppt(path) == [(1, "Grouped heading\nRevenue | 42")]


def test_page_limit_fails_before_inference(tmp_path, monkeypatch):
    path = tmp_path / "large.pdf"
    make_pdf(path)
    monkeypatch.setenv("VLM_MAX_PAGES", "1")
    monkeypatch.setattr(vision.OllamaVisionClient, "describe", lambda *args: pytest.fail("called VLM"))
    with pytest.raises(ValueError, match="1~1"):
        parse_document(path, path.name)


def test_corrupt_pdf_is_not_silently_saved_in_vision_mode(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"invalid")
    with pytest.raises(ValueError, match="손상"):
        parse_document(path, path.name)


def test_missing_converter_is_actionable(tmp_path, monkeypatch):
    path = tmp_path / "slides.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    monkeypatch.setattr(visual_parser.shutil, "which", lambda _: None)
    with pytest.raises(vision.VisionUnavailableError, match="LibreOffice"):
        parse_document(path, path.name)


@pytest.mark.parametrize("url", ["https://ollama.com", "http://192.168.0.2:11434", "http://localhost@evil.test"])
def test_remote_endpoints_are_rejected(url, monkeypatch):
    monkeypatch.setenv("VLM_BASE_URL", url)
    with pytest.raises(ValueError, match="VLM_ALLOWED_HOSTS"):
        vision.OllamaVisionClient()


def test_explicitly_allowlisted_lan_or_docker_endpoint_is_accepted(monkeypatch):
    monkeypatch.setenv("VLM_BASE_URL", "http://192.168.0.20:11434")
    monkeypatch.setenv("VLM_ALLOWED_HOSTS", "192.168.0.20")
    assert vision.OllamaVisionClient().url == "http://192.168.0.20:11434"


def test_auto_mode_keeps_text_parsing_when_vlm_is_not_ready(tmp_path, monkeypatch):
    from app.services import document_service

    source = tmp_path / "slides.pdf"
    source.write_bytes(b"placeholder")
    monkeypatch.setenv("DOCUMENT_VISION_MODE", "auto")
    monkeypatch.setitem(document_service.PARSERS, ".pdf", lambda _: [(1, "텍스트 근거")])
    monkeypatch.setattr(
        vision.OllamaVisionClient,
        "diagnostics",
        lambda self, timeout=1: {"operational": False, "model": self.model},
    )
    result = parse_document(source, source.name)
    assert result.full_text == "텍스트 근거"


def test_cloud_model_rejected(monkeypatch):
    monkeypatch.setenv("VLM_MODEL", "qwen3-vl:235b-cloud")
    with pytest.raises(ValueError, match="클라우드"):
        vision.OllamaVisionClient()


def mock_transport(monkeypatch, handler):
    real_client = httpx.Client
    def client(**kwargs):
        assert kwargs["trust_env"] is False
        assert kwargs["follow_redirects"] is False
        return real_client(transport=httpx.MockTransport(handler), **kwargs)
    monkeypatch.setattr(vision.httpx, "Client", client)


def test_ollama_request_contains_base64_image_and_separate_system_prompt(monkeypatch):
    monkeypatch.setenv("VLM_API_KEY", "test-secret")
    def handler(request):
        body = json.loads(request.content)
        assert request.url.path == "/api/chat"
        assert request.headers["Authorization"] == "Bearer test-secret"
        assert body["stream"] is False
        assert body["model"] == "qwen3-vl:4b-instruct"
        assert body["messages"][0]["role"] == "system"
        assert base64.b64decode(body["messages"][1]["images"][0]) == b"png data"
        return httpx.Response(200, json={"done": True, "message": {"content": "차트"}})
    mock_transport(monkeypatch, handler)
    assert vision.OllamaVisionClient().describe(b"png data", "heading", 1, 2) == "차트"


def test_vision_diagnostics_checks_configured_model_without_exposing_endpoint(monkeypatch):
    mock_transport(monkeypatch, lambda request: httpx.Response(200, json={
        "models": [{"name": "qwen3-vl:4b-instruct"}],
    }))
    assert vision.OllamaVisionClient().diagnostics() == {
        "operational": True,
        "model": "qwen3-vl:4b-instruct",
    }


@pytest.mark.parametrize("payload", [
    {}, {"done": True, "message": {"content": ""}},
    {"done": False, "message": {"content": "partial"}},
    {"done": True, "done_reason": "length", "message": {"content": "truncated"}},
])
def test_invalid_model_response_fails_explicitly(payload, monkeypatch):
    mock_transport(monkeypatch, lambda _: httpx.Response(200, json=payload))
    with pytest.raises(vision.VisionUnavailableError, match="2페이지"):
        vision.OllamaVisionClient().describe(b"png", "", 2, 2)


def test_unavailable_model_and_busy_slot(monkeypatch):
    mock_transport(monkeypatch, lambda _: httpx.Response(404, json={"error": "missing model"}))
    with pytest.raises(vision.VisionUnavailableError):
        vision.OllamaVisionClient().describe(b"png", "", 1, 2)
    visual_parser._VISION_SLOT.acquire()
    try:
        with pytest.raises(vision.VisionUnavailableError, match="다른 문서"):
            visual_parser.enrich_visual_pages(Path("unused"), [])
    finally:
        visual_parser._VISION_SLOT.release()


def test_failed_vision_upload_returns_503_without_persisting(tmp_path, monkeypatch):
    from app.controllers import document_controller
    from app.dependencies import get_current_user, get_document_repository, get_object_storage
    from app.main import app
    from app.repositories.document_repository import InMemoryDocumentRepository
    from tests.test_document_resources import OWNER, FakeStorage

    source = tmp_path / "source.pdf"
    make_pdf(source)
    uploads = tmp_path / "uploads"
    monkeypatch.setattr(document_controller, "UPLOAD_DIR", uploads)
    async def fail(*args):
        raise vision.VisionUnavailableError("시각 분석 모델을 사용할 수 없습니다.")
    monkeypatch.setattr(document_controller, "parse_in_process", fail)
    repository = InMemoryDocumentRepository()
    class NoUpload(FakeStorage):
        async def upload(self, *args):
            pytest.fail("must not save partially analyzed document")
    app.dependency_overrides[get_current_user] = lambda: OWNER
    app.dependency_overrides[get_document_repository] = lambda: repository
    app.dependency_overrides[get_object_storage] = lambda: NoUpload()
    try:
        response = TestClient(app).post("/documents/parse", files={"file": ("source.pdf", source.read_bytes(), "application/pdf")})
        assert response.status_code == 503
        assert asyncio.run(repository.list(OWNER.user_id)) == []
        assert list(uploads.iterdir()) == []
    finally:
        app.dependency_overrides.clear()


@pytest.mark.parametrize("fail_second_page", [False, True])
def test_model_stays_loaded_across_pages_and_unloads_once(tmp_path, monkeypatch, fail_second_page):
    path = tmp_path / "slides.pdf"
    make_pdf(path)
    calls = []
    def handler(request):
        body = json.loads(request.content)
        calls.append((request.url.path, body))
        if request.url.path == "/api/generate":
            return httpx.Response(200, json={"done": True})
        if fail_second_page and len(calls) == 2:
            return httpx.Response(503, json={"error": "failed"})
        return httpx.Response(200, json={"done": True, "message": {"content": "차트"}})
    mock_transport(monkeypatch, handler)
    monkeypatch.setenv("VLM_KEEP_ALIVE_SECONDS", "300")
    monkeypatch.setenv("VLM_UNLOAD_AFTER_DOCUMENT", "true")
    if fail_second_page:
        with pytest.raises(vision.VisionUnavailableError):
            parse_document(path, path.name)
    else:
        parse_document(path, path.name)
    assert [url for url, _ in calls] == ["/api/chat", "/api/chat", "/api/generate"]
    assert [body["keep_alive"] for _, body in calls] == [300, 300, 0]


def test_keep_warm_policy_skips_unload_and_cleanup_failure_does_not_hide_success(monkeypatch):
    calls = []
    def handler(request):
        calls.append(request.url.path)
        if request.url.path == "/api/generate":
            return httpx.Response(503)
        return httpx.Response(200, json={"done": True, "message": {"content": "차트"}})
    mock_transport(monkeypatch, handler)
    monkeypatch.setenv("VLM_UNLOAD_AFTER_DOCUMENT", "false")
    with vision.OllamaVisionClient().document_session() as client:
        assert client.describe(b"png", "", 1, 2) == "차트"
    assert calls == ["/api/chat"]
    monkeypatch.setenv("VLM_UNLOAD_AFTER_DOCUMENT", "true")
    with vision.OllamaVisionClient().document_session() as client:
        assert client.describe(b"png", "", 1, 2) == "차트"
    assert calls[-1] == "/api/generate"
